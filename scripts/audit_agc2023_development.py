#!/usr/bin/env python3
"""Summarize the frozen AGC 2023 pre-trial development archive.

The archive is development-only. This script deliberately creates no holdout,
validation score, or counterfactual-actuator claim.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from statistics import mean, stdev


EXPECTED_FILES = (
    "ClimateTimeseries.xlsx",
    "CropMeasurements.xlsx",
    "DestructiveHarvest.xlsx",
)
CLIMATE_FIELDS = (
    "tout", "rhout", "iglob", "windsp", "rain", "parout", "pyrgeo",
    "co2out", "t_air", "rh", "co2", "scr_enrg", "scr_blck",
    "ligth_on", "vent_lee", "vent_wind", "t_rail", "part1", "par2",
    "par3", "par4",
)
SAMPLE_DATE_DENSITY = {
    "2023-09-04": 48.0,
    "2023-09-18": 48.0,
    "2023-09-29": 25.0,
    "2023-11-09": 20.0,
}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def finite_number(value):
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    if isinstance(value, str):
        try:
            parsed = float(value)
        except ValueError:
            return None
        return parsed if math.isfinite(parsed) else None
    return None


def stats(values: list[float], rows: int) -> dict:
    return {
        "nonmissing": len(values),
        "coverage": len(values) / rows if rows else 0.0,
        "min": min(values) if values else None,
        "max": max(values) if values else None,
        "mean": mean(values) if values else None,
    }


def trusted_sample_density(date: str, workbook_density) -> tuple[float, str | None]:
    density = SAMPLE_DATE_DENSITY[date]
    observed = finite_number(workbook_density)
    if date == "2023-09-04" and observed == 1050.0:
        return density, "combined-table transplant density 1050 conflicts with Info sheet and table geometry; replaced by 48 only for this sampled date"
    if observed != density:
        raise ValueError(f"unexpected density {workbook_density!r} on {date}; expected {density}")
    return density, None


def load_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl is required: install the repository's data-audit extras") from exc
    return load_workbook


def audit_climate(path: Path) -> dict:
    load_workbook = load_openpyxl()
    ws = load_workbook(path, read_only=True, data_only=True)["weather_climate"]
    rows_iter = ws.iter_rows(values_only=True)
    header = list(next(rows_iter))
    index = {name: i for i, name in enumerate(header)}
    missing = [name for name in ("date", *CLIMATE_FIELDS) if name not in index]
    if missing:
        raise ValueError(f"missing climate columns: {missing}")
    values = {name: [] for name in CLIMATE_FIELDS}
    changes = Counter()
    previous = {}
    timestamps = []
    rows = 0
    for row_values in rows_iter:
        timestamp = row_values[index["date"]]
        if not isinstance(timestamp, datetime):
            continue
        rows += 1
        timestamps.append(timestamp)
        for name in CLIMATE_FIELDS:
            value = finite_number(row_values[index[name]])
            if value is None:
                continue
            values[name].append(value)
            if name in previous and previous[name] != value:
                changes[name] += 1
            previous[name] = value
    deltas = [(b - a).total_seconds() for a, b in zip(timestamps, timestamps[1:])]
    return {
        "rows": rows,
        "start": min(timestamps).isoformat() if timestamps else None,
        "end": max(timestamps).isoformat() if timestamps else None,
        "non_five_minute_steps": sum(delta != 300 for delta in deltas),
        "fields": {
            name: {**stats(values[name], rows), "changes": changes[name]}
            for name in CLIMATE_FIELDS
        },
    }


def audit_destructive_crop(path: Path) -> dict:
    load_workbook = load_openpyxl()
    workbook = load_workbook(path, read_only=True, data_only=True)
    ws = workbook["All Data"]
    rows_iter = ws.iter_rows(values_only=True)
    for _ in range(4):
        next(rows_iter)
    header = {value: col for col, value in enumerate(next(rows_iter))}
    required = (
        "Date", "Plant density (p/m2)", "Light", "EC ",
        "Leaf Area measured (cm²/plant)",
    )
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError(f"missing destructive-harvest columns: {missing}")
    groups = defaultdict(list)
    density_pairs = Counter()
    warnings = set()
    rows = 0
    for row_number, row_values in enumerate(rows_iter, start=6):
        date_value = row_values[header["Date"]]
        if not isinstance(date_value, datetime):
            continue
        date = date_value.date().isoformat()
        if date not in SAMPLE_DATE_DENSITY:
            raise ValueError(f"unfrozen destructive-harvest date: {date}")
        workbook_density = row_values[header["Plant density (p/m2)"]]
        density, warning = trusted_sample_density(date, workbook_density)
        if warning:
            warnings.add(warning)
        density_pairs[(date, str(workbook_density))] += 1
        leaf_area = finite_number(row_values[header["Leaf Area measured (cm²/plant)"]])
        if leaf_area is None:
            raise ValueError(f"missing leaf area at row {row_number}")
        light = str(row_values[header["Light"]])
        ec = str(row_values[header["EC "]])
        groups[(date, density, light, ec)].append(leaf_area)
        rows += 1
    summaries = []
    for (date, density, light, ec), leaf_area in sorted(groups.items()):
        summaries.append({
            "date": date,
            "density_plants_m2": density,
            "light": light,
            "ec": ec,
            "n": len(leaf_area),
            "leaf_area_mean_cm2_plant": mean(leaf_area),
            "leaf_area_sd_cm2_plant": stdev(leaf_area) if len(leaf_area) > 1 else 0.0,
            "lai_m2_m2": mean(leaf_area) * density / 10000.0,
        })
    return {
        "rows": rows,
        "groups": summaries,
        "raw_date_density_counts": [
            {"date": date, "workbook_density": density, "rows": count}
            for (date, density), count in sorted(density_pairs.items())
        ],
        "warnings": sorted(warnings),
        "continuous_density_gap": "Info sheet jumps from 22-29 September at 25 plants/m2 to 9 October-9 November at 20 plants/m2; do not infer the 29 September-9 October trajectory from this workbook.",
    }


def audit_nondestructive_crop(path: Path) -> dict:
    load_workbook = load_openpyxl()
    workbook = load_workbook(path, read_only=True, data_only=True)
    ws = workbook["All data"]
    rows_iter = ws.iter_rows(values_only=True)
    header = {value: col for col, value in enumerate(next(rows_iter))}
    required = ("Date", "Treatment", "Variety", "Corrected EC", "Light", "Plantheigth")
    missing = [name for name in required if name not in header]
    if missing:
        raise ValueError(f"missing crop-measurement columns: {missing}")
    dates = Counter()
    rows = 0
    height_rows = 0
    for row_values in rows_iter:
        value = row_values[header["Date"]]
        if not isinstance(value, datetime):
            continue
        rows += 1
        dates[value.date().isoformat()] += 1
        if finite_number(row_values[header["Plantheigth"]]) is not None:
            height_rows += 1
    weekly_2024_typo_sheets = []
    for sheet in ("week 38", "week 39"):
        sheet_ws = workbook[sheet]
        values = sheet_ws.iter_rows(min_row=2, max_col=1, values_only=True)
        years = {row[0].year for row in values if isinstance(row[0], datetime)}
        if 2024 in years:
            weekly_2024_typo_sheets.append(sheet)
    return {
        "rows": rows,
        "dates": dict(sorted(dates.items())),
        "plant_height_nonmissing_rows": height_rows,
        "warning": "Use the consolidated All data sheet. The week 38 and week 39 source tabs contain 2024 date typos; the consolidated sheet contains the internally consistent 2023 dates.",
        "weekly_tabs_with_2024_date_typo": weekly_2024_typo_sheets,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--schema", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    schema = json.loads(args.schema.read_text())
    expected_hashes = schema["workbook_hashes"]
    actual_hashes = {name: digest(args.data_dir / name) for name in EXPECTED_FILES}
    if actual_hashes != expected_hashes:
        raise ValueError(f"workbook hash mismatch: {actual_hashes}")
    result = {
        "audit_id": "agc2023-development-summary-v0",
        "schema_audit_id": schema["audit_id"],
        "data_role": "v9 development only; never prospective validation",
        "workbook_hashes": actual_hashes,
        "climate": audit_climate(args.data_dir / "ClimateTimeseries.xlsx"),
        "destructive_crop": audit_destructive_crop(args.data_dir / "DestructiveHarvest.xlsx"),
        "nondestructive_crop": audit_nondestructive_crop(args.data_dir / "CropMeasurements.xlsx"),
        "claim_boundary": {
            "supported": [
                "develop temperature and RH observed-action replay",
                "develop rail-pipe, ventilation, screens and lighting response",
                "develop dwarf-tomato crop-state and light-treatment components",
            ],
            "unsupported": [
                "prospective validation",
                "CO2 actuator-response identification",
                "fogging-response identification",
                "arbitrary counterfactual action validation",
            ],
        },
    }
    args.out.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n")
    print(json.dumps({
        "rows": result["climate"]["rows"],
        "period": [result["climate"]["start"], result["climate"]["end"]],
        "crop_groups": len(result["destructive_crop"]["groups"]),
        "warnings": result["destructive_crop"]["warnings"] + [result["nondestructive_crop"]["warning"]],
    }, indent=2))


if __name__ == "__main__":
    main()
